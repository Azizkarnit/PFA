import os
import shutil
import io
import datetime
from sqlalchemy.orm import Session
import openpyxl
from typing import Any
import logging

logger = logging.getLogger(__name__)

from app.core.celery_app import celery_app
from app.core.database import SessionLocal
from app.models.survey_questionnaire import SurveyQuestionnaire
from app.models.survey_passage import SurveyPassage
from app.models.sample import Sample
from app.models.sample_company import SampleCompany
from app.models.survey_result_file import SurveyResultFile
from app.models.sector import Sector
from app.models.user import User
from app.models.collection_tracking import CollectionTracking
from app.models.survey import Survey

UPLOAD_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "uploads")
os.makedirs(os.path.join(UPLOAD_DIR, "questionnaires"), exist_ok=True)
os.makedirs(os.path.join(UPLOAD_DIR, "samples"), exist_ok=True)
os.makedirs(os.path.join(UPLOAD_DIR, "results"), exist_ok=True)
EXPORT_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "exports")
os.makedirs(EXPORT_DIR, exist_ok=True)

from app.core.utils import user_display_name

def _user_name(u: User) -> str:
    return user_display_name(u)

def _passage_company_count(passage: SurveyPassage, db: Session) -> int:
    count = 0
    for sample in passage.samples:
        count += db.query(SampleCompany).filter(SampleCompany.sample_id == sample.id).count()
    return count


@celery_app.task(bind=True)
def process_questionnaire_pdf(self, survey_id: int, passage_id: int, temp_file_path: str, original_filename: str):
    db: Session = SessionLocal()
    try:
        p = db.query(SurveyPassage).filter(SurveyPassage.id == passage_id, SurveyPassage.survey_id == survey_id).first()
        if not p:
            raise Exception("Passage not found")

        q = db.query(SurveyQuestionnaire).filter(SurveyQuestionnaire.passage_id == passage_id).order_by(SurveyQuestionnaire.id.desc()).first()
        if not q:
            q = SurveyQuestionnaire(passage_id=passage_id, version="1.0", status="ACTIVE")
            db.add(q)
            db.flush()

        safe_name = f"q_{passage_id}_{original_filename}"
        dest = os.path.join(UPLOAD_DIR, "questionnaires", safe_name)
        
        # Move file
        shutil.move(temp_file_path, dest)
        
        q.questionnaire_pdf_path = safe_name
        db.commit()
        db.refresh(q)

        return {
            "id": q.id, "passage_id": q.passage_id, "questionnaire_url": q.questionnaire_url,
            "questionnaire_pdf_path": q.questionnaire_pdf_path, "version": q.version, 
            "status": q.status, "created_at": q.created_at.isoformat() if q.created_at else None
        }
    except Exception as e:
        if os.path.exists(temp_file_path):
            os.remove(temp_file_path)
        raise e
    finally:
        db.close()


@celery_app.task(bind=True)
def process_sample_upload(self, survey_id: int, passage_id: int, temp_file_path: str, original_filename: str, uploaded_by_id: int):
    db: Session = SessionLocal()
    try:
        p = db.query(SurveyPassage).filter(SurveyPassage.id == passage_id, SurveyPassage.survey_id == survey_id).first()
        if not p:
            raise Exception("Passage not found")

        with open(temp_file_path, "rb") as f:
            content = f.read()

        import csv
        identifiers = set()

        if original_filename.lower().endswith(".csv"):
            content_str = content.decode("utf-8", errors="ignore")
            reader = csv.DictReader(io.StringIO(content_str))
            if reader.fieldnames:
                id_col = next((f for f in reader.fieldnames if f and f.strip().lower() == "identifier"), None)
                if id_col:
                    for row in reader:
                        val = row.get(id_col)
                        if val: identifiers.add(str(val).strip())
        else:
            try:
                wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
                ws = wb.active
                if ws and ws.max_row >= 1:
                    headers = [cell.value for cell in ws[1]]
                    id_idx = next((i for i, h in enumerate(headers) if h and str(h).strip().lower() == "identifier"), None)
                    if id_idx is not None:
                        for row in ws.iter_rows(min_row=2, values_only=True):
                            val = row[id_idx]
                            if val: identifiers.add(str(val).strip())
            except Exception as e:
                logger.warning(f"Error parsing uploaded sample workbook: {e}")

        from app.models.company import Company
        if identifiers:
            companies = db.query(Company).filter(Company.identifier.in_(list(identifiers))).all()
        else:
            companies = []

        row_count = len(companies)

        # Save to permanent storage
        safe_name = f"sample_{passage_id}_{original_filename}"
        dest = os.path.join(UPLOAD_DIR, "samples", safe_name)
        shutil.move(temp_file_path, dest)

        # Delete old
        old = db.query(Sample).filter(Sample.passage_id == passage_id).first()
        if old:
            db.query(SampleCompany).filter(SampleCompany.sample_id == old.id).delete(synchronize_session=False)
            db.delete(old)
            db.flush()
            
        db.query(CollectionTracking).filter(CollectionTracking.passage_id == passage_id).delete(synchronize_session=False)
        db.flush()

        sample = Sample(passage_id=passage_id, uploaded_by=uploaded_by_id, total_companies=row_count)
        db.add(sample)
        db.commit()
        db.refresh(sample)
        
        for c in companies:
            db.add(SampleCompany(sample_id=sample.id, company_id=c.id))
            db.add(CollectionTracking(passage_id=passage_id, company_id=c.id, status="NOT_STARTED"))
        db.commit()

        uploader = db.query(User).filter(User.id == sample.uploaded_by).first()
        return {
            "id": sample.id, "passage_id": sample.passage_id, "total_companies": sample.total_companies,
            "uploaded_by": sample.uploaded_by, "uploader_name": _user_name(uploader) if uploader else None,
            "created_at": sample.created_at.isoformat() if sample.created_at else None
        }
    except Exception as e:
        if os.path.exists(temp_file_path):
            os.remove(temp_file_path)
        raise e
    finally:
        db.close()


@celery_app.task(bind=True)
def process_result_upload(self, survey_id: int, sector_id: int, temp_file_path: str, original_filename: str, uploaded_by_id: int):
    db: Session = SessionLocal()
    try:
        lower = original_filename.lower()
        if lower.endswith(".pdf"):
            ftype = "PDF"
        elif lower.endswith((".xlsx", ".xls")):
            ftype = "EXCEL"
        else:
            raise Exception("Invalid file type")

        safe_name = f"result_{survey_id}_{sector_id}_{original_filename}"
        dest = os.path.join(UPLOAD_DIR, "results", safe_name)
        shutil.move(temp_file_path, dest)

        rf = SurveyResultFile(
            survey_id=survey_id, sector_id=sector_id, file_name=original_filename, 
            file_path=safe_name, file_type=ftype, uploaded_by=uploaded_by_id
        )
        db.add(rf)
        db.commit()
        db.refresh(rf)

        sector = db.query(Sector).filter(Sector.id == sector_id).first()
        uploader = db.query(User).filter(User.id == rf.uploaded_by).first()

        return {
            "id": rf.id, "survey_id": rf.survey_id, "sector_id": rf.sector_id,
            "sector_name": sector.name if sector else None,
            "file_name": rf.file_name, "file_type": rf.file_type,
            "uploaded_by": rf.uploaded_by, "uploader_name": _user_name(uploader) if uploader else None,
            "uploaded_at": rf.uploaded_at.isoformat() if rf.uploaded_at else None
        }
    except Exception as e:
        if os.path.exists(temp_file_path):
            os.remove(temp_file_path)
        raise e
    finally:
        db.close()


@celery_app.task(bind=True)
def generate_survey_report(self, survey_id: int, format_type: str):
    db: Session = SessionLocal()
    try:
        survey = db.query(Survey).filter(Survey.id == survey_id).first()
        if not survey:
            raise Exception("Survey not found")

        passages = db.query(SurveyPassage).filter(SurveyPassage.survey_id == survey_id).order_by(SurveyPassage.year.desc(), SurveyPassage.passage_number.desc()).all()

        if format_type == "excel":
            wb = openpyxl.Workbook()
            ws: Any = wb.active
            ws.title = "Report"
            
            headers = ["Passage Year", "Passage Number", "Status", "Total Companies", "Not Started", "In Progress", "Completed", "Completion %"]
            ws.append(headers)

            passages_any: Any = passages
            for p in passages_any:
                total = _passage_company_count(p, db)
                not_started = db.query(CollectionTracking).filter(CollectionTracking.passage_id == p.id, CollectionTracking.status == "NOT_STARTED").count()
                in_progress = db.query(CollectionTracking).filter(CollectionTracking.passage_id == p.id, CollectionTracking.status == "IN_PROGRESS").count()
                completed = db.query(CollectionTracking).filter(CollectionTracking.passage_id == p.id, CollectionTracking.status == "COMPLETED").count()
                pct = round(completed / total * 100, 1) if total > 0 else 0.0

                ws.append([
                    p.year, p.passage_number, p.status,
                    total, not_started, in_progress, completed, pct
                ])

            filename = f"survey_report_{survey_id}_{datetime.datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
            filepath = os.path.join(EXPORT_DIR, filename)
            wb.save(filepath)  # type: ignore

            return {
                "file_url": f"/api/v1/surveys/downloads/{filename}",
                "filename": filename
            }
        else:
            from reportlab.lib.pagesizes import letter
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib import colors

            filename = f"survey_report_{survey_id}_{datetime.datetime.now().strftime('%Y%m%d%H%M%S')}.pdf"
            filepath = os.path.join(EXPORT_DIR, filename)

            doc = SimpleDocTemplate(filepath, pagesize=letter)
            elements = []
            styles = getSampleStyleSheet()
            elements.append(Paragraph(f"Survey Monitoring Report - {survey.name}", styles['Title']))

            data = [["Year", "Passage", "Status", "Total", "Not Started", "In Progress", "Completed", "Pct %"]]

            passages_any: Any = passages
            for p in passages_any:
                total = _passage_company_count(p, db)
                not_started = db.query(CollectionTracking).filter(CollectionTracking.passage_id == p.id, CollectionTracking.status == "NOT_STARTED").count()
                in_progress = db.query(CollectionTracking).filter(CollectionTracking.passage_id == p.id, CollectionTracking.status == "IN_PROGRESS").count()
                completed = db.query(CollectionTracking).filter(CollectionTracking.passage_id == p.id, CollectionTracking.status == "COMPLETED").count()
                pct = round(completed / total * 100, 1) if total > 0 else 0.0

                data.append([
                    str(p.year), str(p.passage_number), p.status,
                    str(total), str(not_started), str(in_progress), str(completed), str(pct)
                ])

            t = Table(data)
            t.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.grey),
                ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
                ('ALIGN', (0,0), (-1,-1), 'CENTER'),
                ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
                ('BOTTOMPADDING', (0,0), (-1,0), 12),
                ('BACKGROUND', (0,1), (-1,-1), colors.beige),
                ('GRID', (0,0), (-1,-1), 1, colors.black)
            ]))
            elements.append(t)
            doc.build(elements)

            return {
                "file_url": f"/api/v1/surveys/downloads/{filename}",
                "filename": filename
            }
    finally:
        db.close()
