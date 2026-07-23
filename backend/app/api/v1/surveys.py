import os, shutil
from fastapi import APIRouter, Depends, Query, HTTPException, status, UploadFile, File, Request
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session
from sqlalchemy import or_
from typing import Optional, List, Any
from app.core.utils import get_client_ip

from app.core.database import get_db
from app.models.survey import Survey
from app.models.survey_assignment import SurveyAssignment
from app.models.survey_passage import SurveyPassage
from app.models.survey_questionnaire import SurveyQuestionnaire
from app.models.sample import Sample
from app.models.sample_company import SampleCompany
from app.models.collection_tracking import CollectionTracking
from app.models.survey_result_file import SurveyResultFile

from app.models.sector import Sector
from app.models.company import Company
from app.models.user import User
from app.models.role import Role
from app.models.audit_log import AuditLog
from app.schemas.survey import (
    SurveyCreate, SurveyUpdate, SurveyResponse, PaginatedSurveyResponse,
    AssignedAdminResponse,
    PassageCreate, PassageUpdate, PassageResponse, PaginatedPassageResponse,
    QuestionnaireCreate, QuestionnaireResponse,
    SampleResponse,
    MonitoringStats, SurveyMonitoringResponse,
    ResultFileResponse,
)
from app.api.v1.auth import get_current_user
from celery.result import AsyncResult
from app.tasks.survey_tasks import (
    process_questionnaire_pdf, process_sample_upload, 
    process_result_upload, generate_survey_report
)
from app.tasks.email_tasks import send_survey_invitations
import datetime

router = APIRouter()

# ── Upload directory ──────────────────────────────────────────────────────────
UPLOAD_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "..", "uploads")
os.makedirs(os.path.join(UPLOAD_DIR, "questionnaires"), exist_ok=True)
os.makedirs(os.path.join(UPLOAD_DIR, "samples"), exist_ok=True)
os.makedirs(os.path.join(UPLOAD_DIR, "results"), exist_ok=True)
os.makedirs(os.path.join(UPLOAD_DIR, "temp"), exist_ok=True)
EXPORT_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "..", "exports")
os.makedirs(EXPORT_DIR, exist_ok=True)


# ─── Helper: Build SurveyResponse ─────────────────────────────────────────────

def _build_survey_response(survey: Survey, db: Session) -> SurveyResponse:
    s: Any = survey
    admins: List[AssignedAdminResponse] = []
    for assignment in s.assignments:
        u: Any = assignment.user
        if u:
            name = f"{u.first_name or ''} {u.last_name or ''}".strip() or u.email
            admins.append(AssignedAdminResponse(id=u.id, name=name, email=u.email))  # type: ignore

    passages: Any = s.passages
    passage_count = len(passages)
    last_passage = None
    if passages:
        sorted_passages = sorted(passages, key=lambda p: (p.year, p.passage_number), reverse=True)
        last_passage = sorted_passages[0]

    return SurveyResponse(
        id=s.id, code=s.code, name=s.name,
        description=s.description,
        periodicity=s.periodicity,
        status=s.status, assigned_admins=admins,
        passage_count=passage_count,
        last_passage_year=last_passage.year if last_passage else None,
        last_passage_number=last_passage.passage_number if last_passage else None,
        created_by=s.created_by,
        created_at=s.created_at, updated_at=s.updated_at,
    )


from app.core.utils import user_display_name


def _user_name(u: User) -> str:
    return user_display_name(u)


def _passage_company_count(passage: SurveyPassage, db: Session) -> int:
    count = 0
    for sample in passage.samples:
        count += db.query(SampleCompany).filter(SampleCompany.sample_id == sample.id).count()
    return count


# ═══════════════════════════════════════════════════════════════════════════════
# LOOKUPS
# ═══════════════════════════════════════════════════════════════════════════════


@router.get("/admins", response_model=List[AssignedAdminResponse])
def get_survey_admins(db: Session = Depends(get_db)):
    users = (
        db.query(User).join(Role, User.role_id == Role.id)
        .filter(Role.code.in_(["SYSTEM_ADMINISTRATOR", "SURVEY_ADMINISTRATOR"]))
        .filter(User.status == "ACTIVE")
        .order_by(User.first_name, User.last_name).all()
    )
    users_any: Any = users
    return [AssignedAdminResponse(id=u.id, name=_user_name(u), email=u.email) for u in users_any]

@router.get("/assigned-admins-filter", response_model=List[AssignedAdminResponse])
def get_assigned_admins_filter(db: Session = Depends(get_db)):
    users = (
        db.query(User).join(SurveyAssignment, User.id == SurveyAssignment.user_id)
        .filter(User.status == "ACTIVE")
        .distinct()
        .order_by(User.first_name, User.last_name).all()
    )
    users_any: Any = users
    return [AssignedAdminResponse(id=u.id, name=_user_name(u), email=u.email) for u in users_any]


# ═══════════════════════════════════════════════════════════════════════════════
# SURVEYS
# ═══════════════════════════════════════════════════════════════════════════════

@router.get("/", response_model=PaginatedSurveyResponse)
def list_surveys(
    skip: int = Query(0, ge=0), limit: int = Query(10, ge=1, le=100),
    search: Optional[str] = None, status: Optional[str] = None,
    periodicity: Optional[str] = None, assigned_admin_id: Optional[int] = None,
    db: Session = Depends(get_db), current_user: User = Depends(get_current_user),
):
    query = db.query(Survey)
    if current_user.role and current_user.role.code == "SURVEY_ADMINISTRATOR":
        sub = db.query(SurveyAssignment.survey_id).filter(SurveyAssignment.user_id == current_user.id).subquery()
        query = query.filter(Survey.id.in_(sub))  # type: ignore
    if search:
        t = f"%{search}%"
        query = query.filter(or_(Survey.code.ilike(t), Survey.name.ilike(t)))
    if status and status.upper() != "ALL":
        query = query.filter(Survey.status == status.upper())
    if periodicity:
        query = query.filter(Survey.periodicity == periodicity.upper())
    if assigned_admin_id:
        sub = db.query(SurveyAssignment.survey_id).filter(SurveyAssignment.user_id == assigned_admin_id).subquery()
        query = query.filter(Survey.id.in_(sub))  # type: ignore
    total = query.count()
    surveys = query.order_by(Survey.created_at.desc()).offset(skip).limit(limit).all()
    return PaginatedSurveyResponse(total_count=total, items=[_build_survey_response(s, db) for s in surveys])


@router.post("/", response_model=SurveyResponse, status_code=status.HTTP_201_CREATED)
def create_survey(
    request: Request,
    payload: SurveyCreate,
    db: Session = Depends(get_db), current_user: User = Depends(get_current_user),
):
    if current_user.role and current_user.role.code not in ["SYSTEM_ADMINISTRATOR", "SURVEY_ADMINISTRATOR"]:
        raise HTTPException(status_code=403, detail="Only System and Survey Administrators can create surveys.")
    if db.query(Survey).filter(Survey.code == payload.code).first():
        raise HTTPException(status_code=400, detail=f"Survey code '{payload.code}' is already in use.")
    survey = Survey(
        code=payload.code, name=payload.name, description=payload.description,
        periodicity=payload.periodicity, status=payload.status, created_by=current_user.id,
    )
    db.add(survey)
    db.flush()
    # If created by a Survey Administrator, automatically assign them
    if current_user.role and current_user.role.code == "SURVEY_ADMINISTRATOR":
        if current_user.id not in payload.assigned_admin_ids:
            payload.assigned_admin_ids.append(current_user.id)
            
    for admin_id in payload.assigned_admin_ids:
        if not db.query(User).filter(User.id == admin_id).first():
            raise HTTPException(status_code=400, detail=f"Admin user {admin_id} not found.")
        db.add(SurveyAssignment(survey_id=survey.id, user_id=admin_id))
    
    audit_entry = AuditLog(
        ip_address=get_client_ip(request),
        user_id=current_user.id,
        action="CREATE",
        entity_type="SURVEY",
        entity_id=survey.id,
        new_values={"details": f"Created survey '{survey.code}'"}
    )
    db.add(audit_entry)
    
    db.commit()
    db.refresh(survey)
    
    # Notify Survey Administrators assigned to this new survey
    from app.core.notification_service import notify_survey_admins_for_survey
    notify_survey_admins_for_survey(
        db,
        survey.id,
        "New Survey Assignment",
        f"You have been assigned to the new survey '{survey.name}' ({survey.code}).",
        type="INFO",
        action_url=f"/admin/surveys/{survey.id}"
    )
    
    return _build_survey_response(survey, db)


@router.get("/check-code/{code}")
def check_code_available(code: str, db: Session = Depends(get_db)):
    existing = db.query(Survey).filter(Survey.code == code.strip().upper()).first()
    return {"available": existing is None}


@router.get("/{survey_id}", response_model=SurveyResponse)
def get_survey(
    survey_id: int,
    db: Session = Depends(get_db), current_user: User = Depends(get_current_user),
):
    survey = db.query(Survey).filter(Survey.id == survey_id).first()
    if not survey:
        raise HTTPException(status_code=404, detail="Survey not found.")
    if current_user.role and current_user.role.code == "SURVEY_ADMINISTRATOR":
        if not db.query(SurveyAssignment).filter(SurveyAssignment.survey_id == survey_id, SurveyAssignment.user_id == current_user.id).first():
            raise HTTPException(status_code=403, detail="Access denied.")
    return _build_survey_response(survey, db)


@router.patch("/{survey_id}", response_model=SurveyResponse)
def update_survey(
    request: Request,
    survey_id: int, payload: SurveyUpdate,
    db: Session = Depends(get_db), current_user: User = Depends(get_current_user),
):
    survey = db.query(Survey).filter(Survey.id == survey_id).first()
    if not survey:
        raise HTTPException(status_code=404, detail="Survey not found.")
    if payload.name is not None: survey.name = payload.name
    if payload.description is not None: survey.description = payload.description
    if payload.periodicity is not None: survey.periodicity = payload.periodicity
    if payload.status is not None: survey.status = payload.status
    new_admin_ids = set()
    if payload.assigned_admin_ids is not None:
        existing_admin_ids = {a.user_id for a in db.query(SurveyAssignment).filter(SurveyAssignment.survey_id == survey_id).all()}
        db.query(SurveyAssignment).filter(SurveyAssignment.survey_id == survey_id).delete()
        for aid in payload.assigned_admin_ids:
            db.add(SurveyAssignment(survey_id=survey.id, user_id=aid))
        new_admin_ids = set(payload.assigned_admin_ids) - existing_admin_ids
            
    audit_entry = AuditLog(
        ip_address=get_client_ip(request),
        user_id=current_user.id,
        action="UPDATE",
        entity_type="SURVEY",
        entity_id=survey.id,
        new_values={"details": f"Updated survey '{survey.code}'"}
    )
    db.add(audit_entry)
    
    db.commit()
    db.refresh(survey)
    
    if new_admin_ids:
        from app.core.notification_service import create_notification
        for new_admin_id in new_admin_ids:
            create_notification(
                db, 
                new_admin_id, 
                "New Survey Assignment", 
                f"You have been assigned to the survey '{survey.name}' ({survey.code}).", 
                type="INFO", 
                action_url=f"/admin/surveys/{survey.id}"
            )

    return _build_survey_response(survey, db)


@router.patch("/{survey_id}/status", response_model=SurveyResponse)
def update_survey_status(
    request: Request,
    survey_id: int, new_status: str = Query(...),
    db: Session = Depends(get_db), current_user: User = Depends(get_current_user),
):
    survey = db.query(Survey).filter(Survey.id == survey_id).first()
    if not survey:
        raise HTTPException(status_code=404, detail="Survey not found.")
    if new_status.upper() not in {"ACTIVE", "INACTIVE", "ARCHIVED"}:
        raise HTTPException(status_code=400, detail="Invalid status.")
    survey.status = new_status.upper()
    
    audit_entry = AuditLog(
        ip_address=get_client_ip(request),
        user_id=current_user.id,
        action="STATUS_CHANGE",
        entity_type="SURVEY",
        entity_id=survey.id,
        new_values={"details": f"Changed survey '{survey.code}' status to {survey.status}"}
    )
    db.add(audit_entry)
    
    db.commit()
    db.refresh(survey)
    
    # Notify System Admins if survey is archived/inactive
    if survey.status in {"ARCHIVED", "INACTIVE"}:
        from app.core.notification_service import notify_system_admins
        notify_system_admins(
            db,
            f"Survey {survey.status.capitalize()}",
            f"Survey '{survey.name}' ({survey.code}) was marked as {survey.status.capitalize()} by {current_user.first_name}.",
            type="INFO",
            action_url=f"/admin/surveys/{survey.id}"
        )
        
    return _build_survey_response(survey, db)


# ═══════════════════════════════════════════════════════════════════════════════
# PASSAGES
# ═══════════════════════════════════════════════════════════════════════════════

@router.get("/{survey_id}/passages", response_model=PaginatedPassageResponse)
def list_passages(
    survey_id: int,
    skip: int = Query(0, ge=0), limit: int = Query(20, ge=1, le=100),
    db: Session = Depends(get_db), current_user: User = Depends(get_current_user),
):
    if not db.query(Survey).filter(Survey.id == survey_id).first():
        raise HTTPException(status_code=404, detail="Survey not found.")
    query = db.query(SurveyPassage).filter(SurveyPassage.survey_id == survey_id)
    total = query.count()
    passages = query.order_by(SurveyPassage.year.desc(), SurveyPassage.passage_number.desc()).offset(skip).limit(limit).all()
    items = []
    passages_any: Any = passages
    for p in passages_any:
        items.append(PassageResponse(
            id=p.id, survey_id=p.survey_id, year=p.year,
            passage_number=p.passage_number,
            opening_date=p.opening_date, closing_date=p.closing_date,
            status=p.status,
            company_count=_passage_company_count(p, db),
            created_at=p.created_at, updated_at=p.updated_at,
        ))  # type: ignore
    return PaginatedPassageResponse(total_count=total, items=items)


@router.post("/{survey_id}/passages", response_model=PassageResponse, status_code=status.HTTP_201_CREATED)
def create_passage(
    request: Request,
    survey_id: int, payload: PassageCreate,
    db: Session = Depends(get_db), current_user: User = Depends(get_current_user),
):
    survey = db.query(Survey).filter(Survey.id == survey_id).first()
    if not survey:
        raise HTTPException(status_code=404, detail="Survey not found.")
    if survey.status == "ARCHIVED":
        raise HTTPException(status_code=400, detail="Cannot add passages to an archived survey.")
    # Check uniqueness of year + passage_number for this survey
    existing = db.query(SurveyPassage).filter(
        SurveyPassage.survey_id == survey_id,
        SurveyPassage.year == payload.year,
        SurveyPassage.passage_number == payload.passage_number,
    ).first()
    if existing:
        raise HTTPException(status_code=400, detail=f"Passage {payload.passage_number} for year {payload.year} already exists.")
        
    if payload.closing_date.replace(tzinfo=None) <= datetime.datetime.now():
        raise HTTPException(status_code=400, detail="Closing date must be in the future.")
        
    p = SurveyPassage(
        survey_id=survey_id, year=payload.year,
        passage_number=payload.passage_number,
        opening_date=payload.opening_date, closing_date=payload.closing_date,
        status="DRAFT",
    )
    db.add(p)
    db.flush()
    
    audit_entry = AuditLog(
        ip_address=get_client_ip(request),
        user_id=current_user.id,
        action="CREATE",
        entity_type="PASSAGE",
        entity_id=p.id,
        new_values={"details": f"Created passage {p.year}/P{p.passage_number} for survey {survey_id}"}
    )
    db.add(audit_entry)
    
    db.commit()
    db.refresh(p)
    p_any: Any = p
    return PassageResponse(
        id=p_any.id, survey_id=p_any.survey_id, year=p_any.year, passage_number=p_any.passage_number,
        opening_date=p_any.opening_date, closing_date=p_any.closing_date,
        status=p_any.status, company_count=0, created_at=p_any.created_at, updated_at=p_any.updated_at,
    )


@router.patch("/{survey_id}/passages/{passage_id}", response_model=PassageResponse)
def update_passage(
    request: Request,
    survey_id: int, passage_id: int, payload: PassageUpdate,
    db: Session = Depends(get_db), current_user: User = Depends(get_current_user),
):
    p = db.query(SurveyPassage).filter(SurveyPassage.id == passage_id, SurveyPassage.survey_id == survey_id).first()
    if not p:
        raise HTTPException(status_code=404, detail="Passage not found.")
    if payload.opening_date is not None: p.opening_date = payload.opening_date
    if payload.closing_date is not None: 
        if payload.closing_date.replace(tzinfo=None) <= datetime.datetime.now():
            raise HTTPException(status_code=400, detail="Closing date must be in the future.")
        p.closing_date = payload.closing_date
    if payload.status is not None:
        allowed = {"DRAFT", "READY", "OPEN", "CLOSED", "ARCHIVED"}
        if payload.status.upper() not in allowed:
            raise HTTPException(status_code=400, detail=f"Status must be one of {allowed}")
        if payload.status.upper() == "OPEN":
            # Validation: Must have at least 1 sample company
            company_count = _passage_company_count(p, db)
            if company_count == 0:
                raise HTTPException(status_code=400, detail="Cannot open a passage with no sample companies.")
            
            # Validation: Must have a questionnaire
            q = db.query(SurveyQuestionnaire).filter(SurveyQuestionnaire.passage_id == passage_id).order_by(SurveyQuestionnaire.id.desc()).first()
            if not q or (not q.questionnaire_url and not q.questionnaire_pdf_path):
                raise HTTPException(status_code=400, detail="Cannot open a passage without a valid questionnaire (URL or PDF).")
                
            p.activated_by = current_user.id
            p.activated_at = datetime.datetime.now()
        elif payload.status.upper() == "CLOSED":
            p.closed_by = current_user.id
            p.closed_at = datetime.datetime.now()
        p.status = payload.status.upper()
        
    audit_entry = AuditLog(
        ip_address=get_client_ip(request),
        user_id=current_user.id,
        action="UPDATE",
        entity_type="PASSAGE",
        entity_id=p.id,
        new_values={"details": f"Updated passage {p.year}/P{p.passage_number} (Status: {p.status})"}
    )
    db.add(audit_entry)
    
    db.commit()
    db.refresh(p)
    
    if payload.status is not None and payload.status.upper() == "CLOSED":
        from app.core.notification_service import notify_system_admins
        s = db.query(Survey).filter(Survey.id == p.survey_id).first()
        survey_name = s.name if s else p.survey_id
        notify_system_admins(
            db,
            "Passage Closed",
            f"Passage {p.year}/P{p.passage_number} for survey '{survey_name}' was closed. Results are now ready for review.",
            type="SUCCESS",
            action_url=f"/admin/surveys/{p.survey_id}/results"
        )
    
    # Dispatch invitation email task if just opened
    if payload.status is not None and payload.status.upper() == "OPEN":
        # We will dispatch without it and let the task construct it from config.
        send_survey_invitations.delay(p.id)
    p_any: Any = p
    return PassageResponse(
        id=p_any.id, survey_id=p_any.survey_id, year=p_any.year, passage_number=p_any.passage_number,
        opening_date=p_any.opening_date, closing_date=p_any.closing_date,
        status=p_any.status, company_count=_passage_company_count(p, db),
        created_at=p_any.created_at, updated_at=p_any.updated_at,
    )


# ═══════════════════════════════════════════════════════════════════════════════
# QUESTIONNAIRE
# ═══════════════════════════════════════════════════════════════════════════════

@router.get("/{survey_id}/passages/{passage_id}/questionnaire", response_model=Optional[QuestionnaireResponse])
def get_questionnaire(
    survey_id: int, passage_id: int,
    db: Session = Depends(get_db), current_user: User = Depends(get_current_user),
):
    p = db.query(SurveyPassage).filter(SurveyPassage.id == passage_id, SurveyPassage.survey_id == survey_id).first()
    if not p:
        raise HTTPException(status_code=404, detail="Passage not found.")
    q = db.query(SurveyQuestionnaire).filter(SurveyQuestionnaire.passage_id == passage_id).order_by(SurveyQuestionnaire.id.desc()).first()
    if not q:
        return None
    q_any: Any = q
    return QuestionnaireResponse(
        id=q_any.id, passage_id=q_any.passage_id, questionnaire_url=q_any.questionnaire_url,
        questionnaire_pdf_path=q_any.questionnaire_pdf_path,
        version=q_any.version, status=q_any.status, created_at=q_any.created_at,
    )


@router.post("/{survey_id}/passages/{passage_id}/questionnaire", response_model=QuestionnaireResponse)
def upsert_questionnaire(
    survey_id: int, passage_id: int, payload: QuestionnaireCreate,
    db: Session = Depends(get_db), current_user: User = Depends(get_current_user),
):
    p = db.query(SurveyPassage).filter(SurveyPassage.id == passage_id, SurveyPassage.survey_id == survey_id).first()
    if not p:
        raise HTTPException(status_code=404, detail="Passage not found.")
    q = db.query(SurveyQuestionnaire).filter(SurveyQuestionnaire.passage_id == passage_id).order_by(SurveyQuestionnaire.id.desc()).first()
    if q:
        q.questionnaire_url = payload.questionnaire_url
        q.version = payload.version
    else:
        q = SurveyQuestionnaire(
            passage_id=passage_id, questionnaire_url=payload.questionnaire_url,
            version=payload.version, status="ACTIVE",
        )
        db.add(q)
    db.commit()
    db.refresh(q)
    q_any: Any = q
    return QuestionnaireResponse(
        id=q_any.id, passage_id=q_any.passage_id, questionnaire_url=q_any.questionnaire_url,
        questionnaire_pdf_path=q_any.questionnaire_pdf_path,
        version=q_any.version, status=q_any.status, created_at=q_any.created_at,
    )


@router.post("/{survey_id}/passages/{passage_id}/questionnaire/pdf", status_code=status.HTTP_202_ACCEPTED)
async def upload_questionnaire_pdf(
    survey_id: int, passage_id: int,
    file: UploadFile = File(...),
    db: Session = Depends(get_db), current_user: User = Depends(get_current_user),
):
    if not file.filename or not file.filename.lower().endswith(".pdf"):
        raise HTTPException(status_code=400, detail="Only PDF files are accepted.")
    
    p = db.query(SurveyPassage).filter(SurveyPassage.id == passage_id, SurveyPassage.survey_id == survey_id).first()
    if not p:
        raise HTTPException(status_code=404, detail="Passage not found.")

    import uuid
    temp_name = f"temp_{uuid.uuid4()}_{file.filename}"
    temp_path = os.path.join(UPLOAD_DIR, "temp", temp_name)
    content = await file.read()
    with open(temp_path, "wb") as f:
        f.write(content)

    task = process_questionnaire_pdf.delay(survey_id, passage_id, temp_path, file.filename)
    return {"task_id": task.id}


# ═══════════════════════════════════════════════════════════════════════════════
# SAMPLES
# ═══════════════════════════════════════════════════════════════════════════════

@router.get("/samples/template")
def download_sample_template(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    """
    Generate and return a ready-to-fill Excel template for survey samples.
    """
    import openpyxl
    import io
    from fastapi.responses import StreamingResponse

    wb = openpyxl.Workbook()
    ws: Any = wb.active
    ws.title = "Sample Companies"

    # Header
    ws["A1"] = "identifier"
    
    # Optionally add some hints or styles
    from openpyxl.styles import Font, PatternFill
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    
    cell: Any = ws["A1"]
    cell.font = header_font
    cell.fill = header_fill
    ws.column_dimensions["A"].width = 30

    # Add a dummy row for guidance
    from app.models.company import Company
    example_company = db.query(Company).first()
    example_id = example_company.identifier if example_company else "1234567M"
    ws["A2"] = example_id

    buf = io.BytesIO()
    wb.save(buf)  # type: ignore
    buf.seek(0)
    
    return StreamingResponse(
        buf, 
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="survey_sample_template.xlsx"'}
    )

@router.get("/{survey_id}/passages/{passage_id}/sample", response_model=Optional[SampleResponse])
def get_sample(
    survey_id: int, passage_id: int,
    db: Session = Depends(get_db), current_user: User = Depends(get_current_user),
):
    p = db.query(SurveyPassage).filter(SurveyPassage.id == passage_id, SurveyPassage.survey_id == survey_id).first()
    if not p:
        raise HTTPException(status_code=404, detail="Passage not found.")
    sample = db.query(Sample).filter(Sample.passage_id == passage_id).order_by(Sample.id.desc()).first()
    if not sample:
        return None
    uploader = db.query(User).filter(User.id == sample.uploaded_by).first()
    
    # Fetch up to 100 sample companies to prevent huge payloads
    sc_records = db.query(SampleCompany, Company).join(Company, SampleCompany.company_id == Company.id).filter(SampleCompany.sample_id == sample.id).limit(100).all()
    
    companies_info = []
    for sc, comp in sc_records:
        companies_info.append({
            "id": comp.id,
            "identifier": comp.identifier,
            "company_name": comp.company_name,
            "status": comp.status
        })

    sample_any: Any = sample
    return SampleResponse(
        id=sample_any.id, passage_id=sample_any.passage_id,
        total_companies=sample_any.total_companies,
        uploaded_by=sample_any.uploaded_by,
        uploader_name=_user_name(uploader) if uploader else None,
        created_at=sample_any.created_at,
        companies=companies_info
    )


@router.post("/{survey_id}/passages/{passage_id}/sample", status_code=status.HTTP_202_ACCEPTED)
async def upload_sample(
    survey_id: int, passage_id: int,
    file: UploadFile = File(...),
    db: Session = Depends(get_db), current_user: User = Depends(get_current_user),
):
    allowed_ext = (".xlsx", ".xls", ".csv")
    if not file.filename or not any(file.filename.lower().endswith(e) for e in allowed_ext):
        raise HTTPException(status_code=400, detail="Only Excel (.xlsx/.xls) or CSV files are accepted.")
    
    p = db.query(SurveyPassage).filter(SurveyPassage.id == passage_id, SurveyPassage.survey_id == survey_id).first()
    if not p:
        raise HTTPException(status_code=404, detail="Passage not found.")

    import uuid
    temp_name = f"temp_{uuid.uuid4()}_{file.filename}"
    temp_path = os.path.join(UPLOAD_DIR, "temp", temp_name)
    content = await file.read()
    with open(temp_path, "wb") as f:
        f.write(content)

    task = process_sample_upload.delay(survey_id, passage_id, temp_path, file.filename, current_user.id)
    return {"task_id": task.id}


# ═══════════════════════════════════════════════════════════════════════════════
# MONITORING
# ═══════════════════════════════════════════════════════════════════════════════

@router.get("/{survey_id}/monitoring", response_model=SurveyMonitoringResponse)
def get_monitoring(
    survey_id: int,
    db: Session = Depends(get_db), current_user: User = Depends(get_current_user),
):
    survey = db.query(Survey).filter(Survey.id == survey_id).first()
    if not survey:
        raise HTTPException(status_code=404, detail="Survey not found.")

    passages = db.query(SurveyPassage).filter(SurveyPassage.survey_id == survey_id).order_by(SurveyPassage.year.desc(), SurveyPassage.passage_number.desc()).all()

    total_companies_all = 0
    stats_list: List[MonitoringStats] = []

    passages_any: Any = passages
    for p in passages_any:
        total = _passage_company_count(p, db)
        total_companies_all += total

        not_started = db.query(CollectionTracking).filter(
            CollectionTracking.passage_id == p.id,
            CollectionTracking.status == "NOT_STARTED"
        ).count()
        in_progress = db.query(CollectionTracking).filter(
            CollectionTracking.passage_id == p.id,
            CollectionTracking.status == "IN_PROGRESS"
        ).count()
        completed = db.query(CollectionTracking).filter(
            CollectionTracking.passage_id == p.id,
            CollectionTracking.status == "COMPLETED"
        ).count()
        pct = round(completed / total * 100, 1) if total > 0 else 0.0

        stats_list.append(MonitoringStats(
            passage_id=p.id, passage_year=p.year,
            passage_number=p.passage_number, passage_status=p.status,
            total_companies=total, not_started=not_started,
            in_progress=in_progress, completed=completed, completion_pct=pct,
        ))

    return SurveyMonitoringResponse(
        survey_id=survey_id, total_passages=len(passages),
        total_companies_all_passages=total_companies_all,
        passages=stats_list,
    )


# ═══════════════════════════════════════════════════════════════════════════════
# RESULTS
# ═══════════════════════════════════════════════════════════════════════════════

@router.get("/{survey_id}/results")
def list_results(
    survey_id: int,
    skip: int = Query(0, ge=0),
    limit: int = Query(20, ge=1, le=100),
    db: Session = Depends(get_db), current_user: User = Depends(get_current_user),
):
    if not db.query(Survey).filter(Survey.id == survey_id).first():
        raise HTTPException(status_code=404, detail="Survey not found.")
    query = db.query(SurveyResultFile).filter(SurveyResultFile.survey_id == survey_id)
    total = query.count()
    files = query.order_by(SurveyResultFile.uploaded_at.desc()).offset(skip).limit(limit).all()
    result = []
    files_any: Any = files
    for f in files_any:
        sector = db.query(Sector).filter(Sector.id == f.sector_id).first()
        sector_any: Any = sector
        uploader = db.query(User).filter(User.id == f.uploaded_by).first()
        result.append(ResultFileResponse(
            id=f.id, survey_id=f.survey_id, sector_id=f.sector_id,
            sector_name=sector_any.name if sector_any else None,
            file_name=f.file_name, file_path=f.file_path, file_type=f.file_type,
            uploaded_by=f.uploaded_by,
            uploader_name=_user_name(uploader) if uploader else None,
            uploaded_at=f.uploaded_at,
        ))
    return {"total_count": total, "skip": skip, "limit": limit, "items": result}


@router.post("/{survey_id}/results", status_code=status.HTTP_202_ACCEPTED)
async def upload_result(
    survey_id: int,
    sector_id: int = Query(...),
    file: UploadFile = File(...),
    db: Session = Depends(get_db), current_user: User = Depends(get_current_user),
):
    if not db.query(Survey).filter(Survey.id == survey_id).first():
        raise HTTPException(status_code=404, detail="Survey not found.")
    if not db.query(Sector).filter(Sector.id == sector_id).first():
        raise HTTPException(status_code=400, detail="Invalid sector_id.")

    # Validate that this survey actually has companies from this sector in its samples
    company_in_sector = (
        db.query(Company)
        .join(SampleCompany, SampleCompany.company_id == Company.id)
        .join(Sample, Sample.id == SampleCompany.sample_id)
        .join(SurveyPassage, SurveyPassage.id == Sample.passage_id)
        .filter(SurveyPassage.survey_id == survey_id, Company.sector_id == sector_id)
        .first()
    )
    if not company_in_sector:
        raise HTTPException(
            status_code=400, 
            detail="Cannot upload a report for a sector that has no companies participating in this survey."
        )

    lower = (file.filename or "").lower()
    if not (lower.endswith(".pdf") or lower.endswith(".xlsx") or lower.endswith(".xls")):
        raise HTTPException(status_code=400, detail="Only PDF or Excel files are accepted.")

    import uuid
    temp_name = f"temp_{uuid.uuid4()}_{file.filename}"
    temp_path = os.path.join(UPLOAD_DIR, "temp", temp_name)
    content = await file.read()
    with open(temp_path, "wb") as out:
        out.write(content)

    task = process_result_upload.delay(survey_id, sector_id, temp_path, file.filename, current_user.id)
    return {"task_id": task.id}


@router.delete("/{survey_id}/results/{result_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_result(
    survey_id: int, result_id: int,
    db: Session = Depends(get_db), current_user: User = Depends(get_current_user),
):
    rf = db.query(SurveyResultFile).filter(SurveyResultFile.id == result_id, SurveyResultFile.survey_id == survey_id).first()
    if not rf:
        raise HTTPException(status_code=404, detail="Result file not found.")
    rf_any: Any = rf
    dest = os.path.join(UPLOAD_DIR, "results", rf_any.file_path)
    if os.path.exists(dest):
        os.remove(dest)
    db.delete(rf)
    db.commit()

# ═══════════════════════════════════════════════════════════════════════════════
# REPORTS & TASKS
# ═══════════════════════════════════════════════════════════════════════════════

@router.post("/{survey_id}/reports/export", status_code=status.HTTP_202_ACCEPTED)
def export_survey_report(
    survey_id: int, format: str = Query(..., description="excel or pdf"),
    db: Session = Depends(get_db), current_user: User = Depends(get_current_user),
):
    if not db.query(Survey).filter(Survey.id == survey_id).first():
        raise HTTPException(status_code=404, detail="Survey not found.")
    if format not in ["excel", "pdf"]:
        raise HTTPException(status_code=400, detail="Invalid format.")
    
    task = generate_survey_report.delay(survey_id, format)
    return {"task_id": task.id}

@router.get("/tasks/{task_id}")
def get_task_status(task_id: str, current_user: User = Depends(get_current_user)):
    from app.core.celery_app import celery_app
    task_result = celery_app.AsyncResult(task_id)
    result = {
        "task_id": task_id,
        "status": task_result.state,
    }
    if task_result.state == "SUCCESS":
        result["result"] = task_result.result
    elif task_result.state == "FAILURE":
        result["error"] = str(task_result.info)
    return result

@router.get("/downloads/{filename}")
def download_export_file(filename: str, current_user: User = Depends(get_current_user)):
    filepath = os.path.join(EXPORT_DIR, filename)
    if not os.path.exists(filepath):
        raise HTTPException(status_code=404, detail="File not found")
    return FileResponse(filepath)

@router.get("/downloads/questionnaire/{filename}")
def download_questionnaire_file(filename: str):
    filepath = os.path.join(UPLOAD_DIR, "questionnaires", filename)
    if not os.path.exists(filepath):
        raise HTTPException(status_code=404, detail="File not found")
    return FileResponse(filepath, filename=filename)

@router.get("/downloads/result/{filename}")
def download_result_file(filename: str):
    filepath = os.path.join(UPLOAD_DIR, "results", filename)
    if not os.path.exists(filepath):
        raise HTTPException(status_code=404, detail="File not found")
    return FileResponse(filepath, filename=filename)
